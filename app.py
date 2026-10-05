import csv
import html
import io
import json
import mimetypes
import os
import re
import secrets
import sqlite3
import smtplib
import ssl
import time
import zipfile
from datetime import date, datetime, time as datetime_time, timedelta
from email.message import EmailMessage
from email.utils import formatdate
from getpass import getpass
from pathlib import Path, PurePath
from time import time as unix_time

from flask import Flask, abort, flash, g, jsonify, redirect, render_template, request, session, url_for
from flask_login import LoginManager, UserMixin, current_user, login_required, login_user, logout_user
from flask_wtf.csrf import CSRFProtect
from markupsafe import Markup
from openpyxl import load_workbook
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename
from werkzeug.middleware.proxy_fix import ProxyFix

from enviar_correos import (
    CORREO_ORIGEN,
    NOMBRE_REMITENTE,
    PAUSA_ENTRE_ENVIOS,
    SMTP_PORT,
    SMTP_SERVER,
    SMTP_USUARIO,
)


app = Flask(__name__)
APP_ENV = os.environ.get("APP_ENV", "development").lower()
SECRET_KEY = os.environ.get("FLASK_SECRET_KEY")
if APP_ENV == "production":
    required_settings = [
        "FLASK_SECRET_KEY",
        "SMTP_SERVER",
        "SMTP_PORT",
        "SMTP_USUARIO",
        "CORREO_ORIGEN",
        "NOMBRE_REMITENTE",
    ]
    missing_settings = [name for name in required_settings if not os.environ.get(name)]
    if missing_settings:
        raise RuntimeError(
            "Faltan variables de producción: " + ", ".join(missing_settings)
        )
    if len(SECRET_KEY.encode("utf-8")) < 32:
        raise RuntimeError("FLASK_SECRET_KEY debe tener al menos 32 bytes en producción.")
app.config.update(
    SECRET_KEY=SECRET_KEY or secrets.token_hex(32),
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=APP_ENV == "production",
    PERMANENT_SESSION_LIFETIME=timedelta(hours=8),
    MAX_CONTENT_LENGTH=25 * 1024 * 1024,
    DATABASE=str(Path(os.environ.get("APP_DATA_DIR", app.instance_path)) / "correo.sqlite3"),
)
if APP_ENV == "production":
    app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)

REQUIRED_COLUMNS = {"EMPRESA", "DESTINATARIO", "CC"}
MAX_SPREADSHEET_ROWS = int(os.environ.get("MAX_SPREADSHEET_ROWS", "10000"))
MAX_ATTACHMENT_COUNT = 20
JOB_RETENTION_HOURS = 24
ALLOWED_ATTACHMENT_EXTENSIONS = {
    ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".csv", ".txt",
    ".png", ".jpg", ".jpeg", ".gif", ".webp",
}
EMAIL_PATTERN = re.compile(
    r"^[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@"
    r"[A-Za-z0-9](?:[A-Za-z0-9-]{0,61}[A-Za-z0-9])?"
    r"(?:\.[A-Za-z0-9](?:[A-Za-z0-9-]{0,61}[A-Za-z0-9])?)+$"
)

csrf = CSRFProtect(app)
login_manager = LoginManager(app)
login_manager.login_view = "login"
login_manager.login_message = "Inicia sesión para continuar."
login_manager.login_message_category = "error"


class User(UserMixin):
    def __init__(self, user_id: int, username: str):
        self.id = str(user_id)
        self.username = username


def get_db() -> sqlite3.Connection:
    if "db" not in g:
        database_path = Path(app.config["DATABASE"])
        database_path.parent.mkdir(parents=True, exist_ok=True)
        if os.name == "posix":
            os.chmod(database_path.parent, 0o700)
        g.db = sqlite3.connect(database_path, timeout=10)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys=ON")
        g.db.execute("PRAGMA journal_mode=WAL")
        if os.name == "posix" and database_path.exists():
            os.chmod(database_path, 0o600)
    return g.db


@app.teardown_appcontext
def close_db(_error=None):
    database = g.pop("db", None)
    if database is not None:
        database.close()


def init_db() -> None:
    database = get_db()
    database.executescript(
        """
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL UNIQUE COLLATE NOCASE,
            password_hash TEXT NOT NULL,
            active INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );
        CREATE TABLE IF NOT EXISTS login_attempts (
            username TEXT NOT NULL COLLATE NOCASE,
            remote_addr TEXT NOT NULL,
            failures INTEGER NOT NULL,
            window_started REAL NOT NULL,
            locked_until REAL NOT NULL DEFAULT 0,
            PRIMARY KEY (username, remote_addr)
        );
        CREATE TABLE IF NOT EXISTS jobs (
            id TEXT PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id),
            source_name TEXT NOT NULL,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            completed_at TEXT,
            status TEXT NOT NULL,
            total_rows INTEGER NOT NULL,
            sent_count INTEGER NOT NULL DEFAULT 0,
            error_count INTEGER NOT NULL DEFAULT 0,
            dry_run INTEGER NOT NULL DEFAULT 0,
            payload_json TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS job_rows (
            job_id TEXT NOT NULL REFERENCES jobs(id) ON DELETE CASCADE,
            row_number INTEGER NOT NULL,
            recipient TEXT NOT NULL,
            attachment_name TEXT NOT NULL DEFAULT '',
            status TEXT NOT NULL,
            detail TEXT NOT NULL DEFAULT '',
            PRIMARY KEY (job_id, row_number)
        );
        CREATE TABLE IF NOT EXISTS job_attachments (
            job_id TEXT NOT NULL REFERENCES jobs(id) ON DELETE CASCADE,
            lookup_name TEXT NOT NULL,
            filename TEXT NOT NULL,
            mime_type TEXT NOT NULL,
            content BLOB NOT NULL,
            PRIMARY KEY (job_id, lookup_name)
        );
        CREATE INDEX IF NOT EXISTS jobs_user_created ON jobs(user_id, created_at DESC);
        """
    )
    login_columns = {
        row["name"] for row in database.execute("PRAGMA table_info(login_attempts)").fetchall()
    }
    if "remote_addr" not in login_columns:
        database.execute("DROP TABLE login_attempts")
        database.execute(
            """CREATE TABLE login_attempts (
                username TEXT NOT NULL COLLATE NOCASE,
                remote_addr TEXT NOT NULL,
                failures INTEGER NOT NULL,
                window_started REAL NOT NULL,
                locked_until REAL NOT NULL DEFAULT 0,
                PRIMARY KEY (username, remote_addr)
            )"""
        )
    attachment_columns = {
        row["name"] for row in database.execute("PRAGMA table_info(job_attachments)").fetchall()
    }
    if "mime_type" not in attachment_columns:
        database.execute(
            "ALTER TABLE job_attachments ADD COLUMN mime_type TEXT NOT NULL DEFAULT 'application/octet-stream'"
        )
    database.commit()


def purge_expired_data() -> None:
    database = get_db()
    database.execute(
        f"DELETE FROM jobs WHERE created_at < datetime('now', '-{JOB_RETENTION_HOURS} hours')"
    )
    database.execute(
        "DELETE FROM login_attempts WHERE window_started < ? AND locked_until < ?",
        (unix_time() - 1800, unix_time()),
    )
    database.commit()


@login_manager.user_loader
def load_user(user_id: str):
    try:
        database = get_db()
        row = database.execute(
            "SELECT id, username FROM users WHERE id = ? AND active = 1",
            (int(user_id),),
        ).fetchone()
    except (ValueError, sqlite3.Error):
        return None
    return User(row["id"], row["username"]) if row else None


@app.cli.command("create-user")
def create_user_command():
    """Create an application user; the password is entered only in the terminal."""
    init_db()
    username = input("Usuario: ").strip()
    if not re.fullmatch(r"[A-Za-z0-9_.-]{3,64}", username):
        raise SystemExit("Usuario inválido: usa 3-64 letras, números, punto, guion o guion bajo.")
    password = getpass("Contraseña (mínimo 12 caracteres): ")
    confirmation = getpass("Repite la contraseña: ")
    if len(password) < 12 or password != confirmation:
        raise SystemExit("Las contraseñas no coinciden o no alcanzan 12 caracteres.")
    try:
        get_db().execute(
            "INSERT INTO users (username, password_hash) VALUES (?, ?)",
            (username, generate_password_hash(password)),
        )
        get_db().commit()
    except sqlite3.IntegrityError as error:
        raise SystemExit("Ese usuario ya existe.") from error
    finally:
        password = ""
        confirmation = ""
    print(f"Usuario {username} creado.")


@app.cli.command("init-db")
def init_db_command():
    """Initialize the SQLite database schema."""
    init_db()
    print(f"Base de datos inicializada en {app.config['DATABASE']}.")


def login_locked(username: str, remote_addr: str) -> bool:
    row = get_db().execute(
        "SELECT locked_until FROM login_attempts WHERE username = ? AND remote_addr = ?",
        (username, remote_addr),
    ).fetchone()
    return bool(row and row["locked_until"] > unix_time())


def record_login_failure(username: str, remote_addr: str) -> None:
    now = unix_time()
    database = get_db()
    row = database.execute(
        "SELECT failures, window_started, locked_until FROM login_attempts WHERE username = ? AND remote_addr = ?",
        (username, remote_addr),
    ).fetchone()
    if row is None or now - row["window_started"] > 900:
        failures = 1
        window_started = now
        locked_until = 0
    elif row["locked_until"] > now:
        return
    else:
        failures = row["failures"] + 1
        window_started = row["window_started"]
        locked_until = now + 900 if failures >= 5 else 0
    database.execute(
          """INSERT INTO login_attempts (username, remote_addr, failures, window_started, locked_until)
              VALUES (?, ?, ?, ?, ?)
           ON CONFLICT(username, remote_addr) DO UPDATE SET failures=excluded.failures,
              window_started=excluded.window_started, locked_until=excluded.locked_until""",
          (username, remote_addr, failures, window_started, locked_until),
    )
    database.commit()
DEFAULT_BODY = """<!doctype html>
<html lang="es"><body style="margin:0;font-family:Arial,sans-serif;color:#222;">
<div style="max-width:900px;padding:12px 0;line-height:1.45;">
<p>Estimado Cliente:</p>
<p>Como parte de nuestro compromiso con la continuidad y estabilidad de los servicios, les informamos que se llevarán a cabo trabajos de mantenimiento correctivo de emergencia, correspondientes al {{MOTIVO}}.</p>
<p>Los trabajos se encuentran programados de acuerdo con las siguientes ventanas:</p>
{{#VENTANAS}}{{VENTANAS}}{{/VENTANAS}}
{{#DURACION}}<p><strong>Horas de duración:</strong> {{DURACION}}</p>{{/DURACION}}
{{#SEDE}}<p><strong>SEDE:</strong><br>{{SEDE}}</p>{{/SEDE}}
{{#NC}}<p><strong>NC:</strong><br>{{NC}}</p>{{/NC}}
<p>Durante el tiempo de ejecución de los trabajos podrían presentarse posibles cortes e intermitencias en el servicio dentro del periodo comprendido de la actividad.</p>
<p>Lamentamos profundamente las molestias ocasionadas, asimismo, le indicamos que nos encontramos tomando las medidas necesarias para asegurar su realización de forma satisfactoria, minimizando el impacto en la continuidad de sus servicios.</p>
<p>Agradecemos su comprensión.</p>
<p>Saludos cordiales</p>
<p><strong><span style="color:#000080;font-size:12px;">Plataforma de Atención Clientes (PAC)</span></strong></p>
<p style="color:#000080;font-size:12px;">Gtd Perú<br>Av. La Encalada 1257, piso 14, Santiago de Surco, Lima, Perú<br>Tel: (51 1) 743-0911 / 0800-7-0202<br><img src="https://www.gtd.cl/images/079ff2755e4e1362b5e96035e0f75d4e/original.jpg" alt="Gtd Perú" style="max-width:480px;height:auto;"></p>
<p><strong>Estimado Cliente:</strong> Tener en cuenta que posterior a las 24 horas de haber solicitado su validación, al no tener respuesta por este medio se procederá al cierre administrativo del caso.</p>
</div></body></html>"""
DEFAULT_TEXT_BODY = """Estimado Cliente:

Como parte de nuestro compromiso con la continuidad y estabilidad de los servicios, les informamos que se llevarán a cabo trabajos de mantenimiento correctivo de emergencia, correspondientes al {{MOTIVO}}.

Los trabajos se encuentran programados de acuerdo con las siguientes ventanas:
{{VENTANAS}}
{{#DURACION}}Horas de duración: {{DURACION}}{{/DURACION}}
{{#SEDE}}SEDE:
{{SEDE}}{{/SEDE}}
{{#NC}}NC:
{{NC}}{{/NC}}

Durante el tiempo de ejecución de los trabajos podrían presentarse posibles cortes e intermitencias en el servicio dentro del periodo comprendido de la actividad.

Lamentamos profundamente las molestias ocasionadas, asimismo, le indicamos que nos encontramos tomando las medidas necesarias para asegurar su realización de forma satisfactoria, minimizando el impacto en la continuidad de sus servicios.

Agradecemos su comprensión.

Saludos cordiales"""
FIXED_SIGNATURE_HTML = """<p><strong><span style="color:#000080;font-size:12px;">Plataforma de Atención Clientes (PAC)</span></strong></p>
<p style="color:#000080;font-size:12px;">Gtd Perú<br>Av. La Encalada 1257, piso 14, Santiago de Surco, Lima, Perú<br>Tel: (51 1) 743-0911 / 0800-7-0202<br><img src="https://www.gtd.cl/images/079ff2755e4e1362b5e96035e0f75d4e/original.jpg" alt="Gtd Perú" style="max-width:480px;height:auto;"></p>
<p><strong>Estimado Cliente:</strong> Tener en cuenta que posterior a las 24 horas de haber solicitado su validación, al no tener respuesta por este medio se procederá al cierre administrativo del caso.</p>"""


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.time() == datetime_time.min:
            return value.strftime("%d/%m/%Y")
        return value.strftime("%d/%m/%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, datetime_time):
        return value.strftime("%H:%M")
    return str(value).strip()


WEEKDAYS_ES = (
    "lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"
)


def parse_date(value: str):
    cleaned = (value or "").strip()
    for date_format in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(cleaned, date_format).date()
        except ValueError:
            continue
    match = re.fullmatch(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", cleaned)
    if not match:
        return None
    first, second, year = (int(part) for part in match.groups())
    if first > 12:
        day, month = first, second
    elif second > 12:
        month, day = first, second
    else:
        day, month = first, second
    try:
        return date(year, month, day)
    except ValueError:
        return None


def parse_time(value: str):
    cleaned = (value or "").strip().lower().replace(" ", "").replace(".", "")
    for time_format in ("%H:%M", "%H:%M:%S", "%I:%M%p", "%I:%M:%S%p"):
        try:
            return datetime.strptime(cleaned, time_format).time()
        except ValueError:
            continue
    return None


def format_date_value(value: str) -> str:
    parsed = parse_date(value)
    if parsed is None:
        return value
    weekday = WEEKDAYS_ES[parsed.weekday()].capitalize()
    return f"{weekday} {parsed:%d/%m/%Y}"


def format_time_value(value: str) -> str:
    parsed = parse_time(value)
    if parsed is None:
        return value
    period = "a.m." if parsed.hour < 12 else "p.m."
    hour = parsed.hour % 12 or 12
    return f"{hour:02d}:{parsed.minute:02d} {period}"


def format_duration(value: str) -> str:
    cleaned = (value or "").strip()
    match = re.fullmatch(r"(\d{1,3}):(\d{2})", cleaned)
    if match:
        hours, minutes = (int(part) for part in match.groups())
        parts = []
        if hours:
            parts.append(f"{hours} {'hora' if hours == 1 else 'horas'}")
        if minutes:
            parts.append(f"{minutes} {'minuto' if minutes == 1 else 'minutos'}")
        return " ".join(parts)
    match = re.fullmatch(r"(\d+)\s*(?:h|horas?)?", cleaned, flags=re.IGNORECASE)
    if match:
        hours = int(match.group(1))
        return f"{hours} {'hora' if hours == 1 else 'horas'}"
    return cleaned


def build_windows(row: dict, html_mode: bool = False) -> str:
    start_date = parse_date(row.get("FECHA_INICIO", ""))
    end_date = parse_date(row.get("FECHA_TERMINO", ""))
    start_time = parse_time(row.get("HORA_INICIO", ""))
    end_time = parse_time(row.get("HORA_CULMINACION", ""))
    if None in (start_date, end_date, start_time, end_time) or end_date < start_date:
        return ""

    window_count = max(1, (end_date - start_date).days)
    lines = []
    for offset in range(window_count):
        window_start = start_date + timedelta(days=offset)
        window_end = window_start + timedelta(days=1)
        start_label = format_date_value(window_start.isoformat())
        end_label = format_date_value(window_end.isoformat())
        lines.append(
            f"{start_label} desde las {format_time_value(start_time.strftime('%H:%M'))} "
            f"hasta el {end_label} a las {format_time_value(end_time.strftime('%H:%M'))}"
        )
    if html_mode:
        return "".join(f"<p>{html.escape(line)}</p>" for line in lines)
    return "\n".join(lines)


def _table_rows(headers, values):
    normalized = [str(header or "").strip().lstrip("\ufeff").upper() for header in headers]
    if len(normalized) != len(set(normalized)):
        raise ValueError("El archivo contiene encabezados repetidos.")
    missing = REQUIRED_COLUMNS - set(normalized)
    if missing:
        raise ValueError(f"Faltan columnas obligatorias: {', '.join(sorted(missing))}.")
    rows = []
    for row_number, cells in enumerate(values, start=2):
        if row_number > MAX_SPREADSHEET_ROWS + 1:
            raise ValueError(f"El archivo supera el máximo de {MAX_SPREADSHEET_ROWS} filas de datos.")
        row = {
            header: _cell_text(cells[index] if index < len(cells) else "")
            for index, header in enumerate(normalized)
        }
        if not any(value.strip() for value in row.values()):
            continue
        for optional_column in (
            "ARCHIVO", "SEDE", "NC", "FECHA", "MOTIVO", "FECHA_INICIO",
            "HORA_INICIO", "FECHA_TERMINO", "HORA_CULMINACION", "DURACION",
        ):
            row.setdefault(optional_column, "")
        if "MOTIVO" in row and not row["MOTIVO"]:
            row["MOTIVO"] = "cambio de una mufa crítica de la infraestructura externa"
        row["_row_number"] = row_number
        rows.append(row)
    if not rows:
        raise ValueError("El archivo no contiene filas con datos.")
    return normalized, rows


def read_spreadsheet(filename: str, content: bytes):
    extension = PurePath(filename).suffix.lower()
    if extension == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                expanded_size = sum(item.file_size for item in archive.infolist())
                if expanded_size > 100 * 1024 * 1024 or "xl/workbook.xml" not in archive.namelist():
                    raise ValueError("El Excel no es válido o supera el límite interno de descompresión.")
        except zipfile.BadZipFile as error:
            raise ValueError("El archivo .xlsx no es un libro Excel válido.") from error
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            iterator = workbook.active.iter_rows(values_only=True)
            try:
                headers = next(iterator)
            except StopIteration as error:
                raise ValueError("La hoja de Excel está vacía.") from error
            return _table_rows(headers, iterator)
        finally:
            workbook.close()
    if extension != ".csv":
        raise ValueError("Selecciona un archivo CSV o Excel .xlsx.")

    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1252")
    header_line = text.splitlines()[0] if text.splitlines() else ""
    delimiter = max(",;\t", key=lambda candidate: len(next(csv.reader([header_line], delimiter=candidate))))
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    try:
        headers = next(reader)
    except StopIteration as error:
        raise ValueError("El archivo CSV está vacío.") from error
    return _table_rows(headers, reader)


def validate_attachment(filename: str, content: bytes) -> dict:
    safe_name = secure_filename(filename)
    suffix = Path(safe_name).suffix.lower()
    if not safe_name or suffix not in ALLOWED_ATTACHMENT_EXTENSIONS:
        raise ValueError(
            "Tipo de adjunto no permitido. Formatos aceptados: PDF, DOC/DOCX, XLS/XLSX, CSV, TXT, PNG, JPG, GIF y WEBP."
        )
    signatures = {
        ".pdf": content.startswith(b"%PDF-"),
        ".png": content.startswith(b"\x89PNG\r\n\x1a\n"),
        ".jpg": content.startswith(b"\xff\xd8\xff"),
        ".jpeg": content.startswith(b"\xff\xd8\xff"),
        ".gif": content.startswith((b"GIF87a", b"GIF89a")),
        ".webp": len(content) >= 12 and content[:4] == b"RIFF" and content[8:12] == b"WEBP",
        ".doc": content.startswith(bytes.fromhex("D0CF11E0A1B11AE1")),
        ".xls": content.startswith(bytes.fromhex("D0CF11E0A1B11AE1")),
    }
    if suffix in signatures and not signatures[suffix]:
        raise ValueError(f"El contenido no coincide con la extensión del adjunto '{safe_name}'.")
    if suffix in {".docx", ".xlsx"}:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                expanded_size = sum(item.file_size for item in archive.infolist())
                expected_member = "word/document.xml" if suffix == ".docx" else "xl/workbook.xml"
                if expanded_size > 100 * 1024 * 1024 or expected_member not in archive.namelist():
                    raise ValueError(f"El archivo '{safe_name}' no es un documento Office válido o supera el límite interno.")
        except zipfile.BadZipFile as error:
            raise ValueError(f"El adjunto '{safe_name}' no es un archivo Office válido.") from error
    if suffix in {".csv", ".txt"} and b"\x00" in content:
        raise ValueError(f"El adjunto de texto '{safe_name}' contiene bytes nulos.")
    mime_type = mimetypes.guess_type(safe_name)[0] or "application/octet-stream"
    return {"filename": safe_name, "content": content, "mime_type": mime_type}


def split_emails(value: str) -> list[str]:
    return [address.strip() for address in re.split(r"[;,]", value or "") if address.strip()]


def valid_email(address: str) -> bool:
    if not EMAIL_PATTERN.fullmatch(address):
        return False
    local, domain = address.rsplit("@", 1)
    return not local.startswith(".") and not local.endswith(".") and ".." not in local and ".." not in domain


def validate_rows(rows, attach_enabled: bool, attachment_names: set[str], body_template: str = ""):
    problems = []
    for row in rows:
        row_errors = []
        if not row.get("EMPRESA", "").strip():
            row_errors.append("EMPRESA está vacío")
        for column in ("DESTINATARIO", "CC"):
            addresses = split_emails(row.get(column, ""))
            if not addresses:
                row_errors.append(f"{column} es obligatorio")
            invalid = [address for address in addresses if not valid_email(address)]
            if invalid:
                row_errors.append(f"{column} contiene correo(s) inválido(s): {', '.join(invalid)}")
        filename = row.get("ARCHIVO", "").strip()
        if attach_enabled and filename:
            safe_name = secure_filename(filename.replace("\\", "/").split("/")[-1])
            if safe_name.casefold() not in attachment_names:
                row_errors.append(f"No se cargó el adjunto '{filename}'")
        window_fields = ("FECHA_INICIO", "HORA_INICIO", "FECHA_TERMINO", "HORA_CULMINACION")
        if "{{VENTANAS}}" in body_template and any(row.get(field, "").strip() for field in window_fields):
            missing_window_fields = [field for field in window_fields if not row.get(field, "").strip()]
            if missing_window_fields:
                row_errors.append(f"Faltan datos de la ventana: {', '.join(missing_window_fields)}")
            else:
                start_date = parse_date(row["FECHA_INICIO"])
                end_date = parse_date(row["FECHA_TERMINO"])
                if start_date is None or end_date is None:
                    row_errors.append("FECHA_INICIO o FECHA_TERMINO no tiene un formato de fecha válido")
                elif end_date < start_date:
                    row_errors.append("FECHA_TERMINO debe ser igual o posterior a FECHA_INICIO")
                if parse_time(row["HORA_INICIO"]) is None or parse_time(row["HORA_CULMINACION"]) is None:
                    row_errors.append("HORA_INICIO o HORA_CULMINACION no tiene un formato de hora válido")
        if row_errors:
            problems.append({"row": row["_row_number"], "errors": row_errors})
    return problems


def render_body(template: str, row: dict, html_mode: bool = True) -> str:
    rendered = template
    template_values = dict(row)
    template_values["VENTANAS"] = build_windows(row, html_mode)
    for name, block in re.findall(r"{{#([A-Z_]+)}}(.*?){{/\1}}", rendered, flags=re.DOTALL):
        rendered = rendered.replace(
            f"{{{{#{name}}}}}{block}{{{{/{name}}}}}",
            block if template_values.get(name, "").strip() else "",
        )
    for name, value in template_values.items():
        if name.startswith("_"):
            continue
        if name in {"FECHA", "FECHA_INICIO", "FECHA_TERMINO"}:
            value = format_date_value(value)
        elif name in {"HORA_INICIO", "HORA_CULMINACION"}:
            value = format_time_value(value)
        elif name == "DURACION":
            value = format_duration(value)
        if name == "VENTANAS":
            rendered = rendered.replace("{{VENTANAS}}", value)
            continue
        if html_mode:
            escaped = "<br>".join(html.escape(part.strip()) for part in value.split(";"))
        else:
            escaped = "\n".join(part.strip() for part in value.split(";"))
        rendered = rendered.replace("{{" + name + "}}", escaped)
    return rendered


def make_subject(row: dict, prefix: str, suffix_mode: str, suffix_text: str) -> str:
    suffix = format_date_value(row.get("FECHA", "")) if suffix_mode == "fecha" else suffix_text
    parts = [prefix.strip(), row.get("EMPRESA", "").strip(), suffix.strip()]
    return " | ".join(part.replace("\r", " ").replace("\n", " ") for part in parts if part)


def make_message(row, campaign, sender_name: str, sender_email: str) -> EmailMessage:
    message = EmailMessage()
    message["From"] = f"{sender_name} <{sender_email}>"
    message["To"] = ", ".join(split_emails(row["DESTINATARIO"]))
    message["Cc"] = ", ".join(split_emails(row["CC"]))
    message["Date"] = formatdate(localtime=True)
    message["Subject"] = make_subject(
        row, campaign["subject_prefix"], campaign["suffix_mode"], campaign["suffix_text"]
    )
    body = render_body(campaign["body_template"], row, False)
    plain_signature = (
        "\n\nPlataforma de Atención Clientes (PAC)\nGtd Perú\n"
        "Av. La Encalada 1257, piso 14, Santiago de Surco, Lima, Perú\n"
        "Tel: (51 1) 743-0911 / 0800-7-0202\n\n"
        "Estimado Cliente: Tener en cuenta que posterior a las 24 horas de haber solicitado su validación, "
        "al no tener respuesta por este medio se procederá al cierre administrativo del caso."
    )
    message.set_content(body.rstrip() + plain_signature)
    html_body = html.escape(body).replace("\r\n", "\n").replace("\n", "<br>\n")
    message.add_alternative(
        f"<!doctype html><html lang=\"es\"><body style=\"font-family:Arial,sans-serif;color:#222;line-height:1.45;\">"
        f"<div>{html_body}</div>{FIXED_SIGNATURE_HTML}</body></html>",
        subtype="html",
    )
    attachment_name = row.get("ARCHIVO", "").strip()
    if campaign["attach_enabled"] and attachment_name:
        lookup = secure_filename(attachment_name.replace("\\", "/").split("/")[-1]).casefold()
        attachment = campaign["attachments"][lookup]
        mime_type = attachment.get("mime_type") or mimetypes.guess_type(attachment["filename"])[0] or "application/octet-stream"
        maintype, subtype = mime_type.split("/", 1)
        message.add_attachment(
            attachment["content"],
            maintype=maintype,
            subtype=subtype,
            filename=attachment["filename"],
        )
    return message


def send_campaign(campaign, password: str):
    server = smtplib.SMTP(SMTP_SERVER, SMTP_PORT, timeout=30)
    results = []
    try:
        server.ehlo()
        server.starttls(context=ssl.create_default_context())
        server.ehlo()
        server.login(SMTP_USUARIO, password)
        for index, row in enumerate(campaign["rows"], start=1):
            try:
                message = make_message(row, campaign, NOMBRE_REMITENTE, CORREO_ORIGEN)
                envelope = split_emails(row["DESTINATARIO"]) + split_emails(row["CC"])
                rejected = server.send_message(message, to_addrs=envelope)
                results.append({
                    "row": row["_row_number"],
                    "recipient": row["DESTINATARIO"],
                    "attachment": row.get("ARCHIVO", ""),
                    "ok": not rejected,
                    "detail": "Enviado" if not rejected else "Uno o más destinatarios fueron rechazados.",
                })
            except Exception as error:
                results.append({
                    "row": row["_row_number"],
                    "recipient": row["DESTINATARIO"],
                    "attachment": row.get("ARCHIVO", ""),
                    "ok": False,
                    "detail": f"Fallo SMTP ({type(error).__name__}); revisar con el administrador.",
                })
            if index < len(campaign["rows"]):
                time.sleep(PAUSA_ENTRE_ENVIOS)
    finally:
        server.quit()
    return results


def persist_campaign(campaign_id: str, campaign: dict, user_id: str) -> None:
    database = get_db()
    purge_expired_data()
    payload = {
        key: campaign[key]
        for key in (
            "rows", "headers", "source_name", "attach_enabled", "subject_prefix",
            "suffix_mode", "suffix_text", "body_template", "body_format", "problems",
        )
    }
    status = "invalid" if campaign["problems"] else "ready"
    database.execute(
        """INSERT INTO jobs (id, user_id, source_name, status, total_rows, dry_run, payload_json)
           VALUES (?, ?, ?, ?, ?, 0, ?)""",
        (
            campaign_id,
            int(user_id),
            campaign["source_name"],
            status,
            len(campaign["rows"]),
            json.dumps(payload, ensure_ascii=False),
        ),
    )
    problem_rows = {problem["row"]: "; ".join(problem["errors"]) for problem in campaign["problems"]}
    for row in campaign["rows"]:
        database.execute(
            """INSERT INTO job_rows (job_id, row_number, recipient, attachment_name, status, detail)
               VALUES (?, ?, ?, ?, ?, ?)""",
            (
                campaign_id,
                row["_row_number"],
                "; ".join(filter(None, [row.get("DESTINATARIO", ""), row.get("CC", "")])),
                row.get("ARCHIVO", ""),
                "invalid" if row["_row_number"] in problem_rows else "ready",
                problem_rows.get(row["_row_number"], ""),
            ),
        )
    for key, attachment in campaign["attachments"].items():
        database.execute(
            "INSERT INTO job_attachments (job_id, lookup_name, filename, mime_type, content) VALUES (?, ?, ?, ?, ?)",
            (campaign_id, key, attachment["filename"], attachment["mime_type"], attachment["content"]),
        )
    database.commit()


def load_campaign(campaign_id: str, user_id: str):
    database = get_db()
    job = database.execute(
        "SELECT * FROM jobs WHERE id = ? AND user_id = ?",
        (campaign_id, int(user_id)),
    ).fetchone()
    if not job:
        return None, None
    campaign = json.loads(job["payload_json"])
    campaign["owner_id"] = str(job["user_id"])
    campaign["attachments"] = {
        attachment["lookup_name"]: {
            "filename": attachment["filename"],
            "mime_type": attachment["mime_type"],
            "content": attachment["content"],
        }
        for attachment in database.execute(
            "SELECT lookup_name, filename, mime_type, content FROM job_attachments WHERE job_id = ?",
            (campaign_id,),
        ).fetchall()
    }
    campaign["job_status"] = job["status"]
    return campaign, job


def finish_campaign(campaign_id: str, results: list[dict]) -> None:
    database = get_db()
    success_count = sum(1 for result in results if result["ok"])
    error_count = len(results) - success_count
    for result in results:
        database.execute(
            """UPDATE job_rows SET status = ?, detail = ?
               WHERE job_id = ? AND row_number = ?""",
            (
                "sent" if result["ok"] else "failed",
                result.get("detail", "")[:500],
                campaign_id,
                result["row"],
            ),
        )
    job_status = "sent" if error_count == 0 else ("failed" if success_count == 0 else "partial_failure")
    database.execute(
        """UPDATE jobs SET status = ?, dry_run = ?, sent_count = ?, error_count = ?,
           payload_json = '{}', source_name = 'Lote', completed_at = CURRENT_TIMESTAMP
           WHERE id = ?""",
        (job_status, 0, success_count, error_count, campaign_id),
    )
    database.execute(
        "UPDATE job_rows SET recipient = '', attachment_name = '', detail = '' WHERE job_id = ?",
        (campaign_id,),
    )
    database.execute("DELETE FROM job_attachments WHERE job_id = ?", (campaign_id,))
    database.commit()


def fail_campaign(campaign_id: str, detail: str = "") -> None:
    failure_detail = detail or "El procesamiento del lote falló; revisar el servicio antes de repetir."
    database = get_db()
    database.execute(
        "UPDATE job_rows SET status = 'failed', detail = ? WHERE job_id = ? AND status IN ('ready', 'queued', 'processing')",
        (failure_detail[:500], campaign_id),
    )
    database.execute(
        """UPDATE jobs SET status = 'failed', error_count = total_rows, payload_json = '{}',
           source_name = 'Lote',
           completed_at = CURRENT_TIMESTAMP WHERE id = ?""",
        (campaign_id,),
    )
    database.execute(
        "UPDATE job_rows SET recipient = '', attachment_name = '' WHERE job_id = ?",
        (campaign_id,),
    )
    database.execute("DELETE FROM job_attachments WHERE job_id = ?", (campaign_id,))
    database.commit()


@app.context_processor
def template_helpers():
    return {
        "make_subject": make_subject,
        "render_body": render_body,
        "smtp_user": SMTP_USUARIO,
        "smtp_server": SMTP_SERVER,
        "smtp_port": SMTP_PORT,
        "sender": CORREO_ORIGEN,
        "default_body": DEFAULT_BODY,
        "default_text_body": DEFAULT_TEXT_BODY,
        "fixed_signature": Markup(FIXED_SIGNATURE_HTML),
    }


@app.route("/login", methods=["GET", "POST"])
def login():
    init_db()
    if current_user.is_authenticated:
        return redirect(url_for("index"))
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        remote_addr = request.remote_addr or "unknown"
        locked = login_locked(username, remote_addr)
        if len(username) > 64 or len(password) > 1024 or locked:
            flash("Usuario o contraseña incorrectos, o cuenta temporalmente bloqueada.", "error")
            return render_template("login.html"), 429 if locked else 401
        user_row = get_db().execute(
            "SELECT id, username, password_hash FROM users WHERE username = ? AND active = 1",
            (username,),
        ).fetchone()
        if user_row and check_password_hash(user_row["password_hash"], password):
            get_db().execute(
                "DELETE FROM login_attempts WHERE username = ? AND remote_addr = ?",
                (username, remote_addr),
            )
            get_db().commit()
            session.clear()
            login_user(User(user_row["id"], user_row["username"]), fresh=True)
            session.permanent = True
            return redirect(url_for("index"))
        record_login_failure(username or "<empty>", remote_addr)
        flash("Usuario o contraseña incorrectos, o cuenta temporalmente bloqueada.", "error")
        return render_template("login.html"), 401
    return render_template("login.html")


@app.post("/logout")
@login_required
def logout():
    logout_user()
    return redirect(url_for("login"))


@app.get("/healthz")
def healthz():
    try:
        get_db().execute("SELECT 1").fetchone()
    except sqlite3.Error:
        return jsonify(status="unavailable"), 503
    return jsonify(status="ok")


@app.get("/")
@login_required
def index():
    return render_template(
        "index.html",
        smtp_server=SMTP_SERVER,
        smtp_port=SMTP_PORT,
        sender=CORREO_ORIGEN,
        default_body=DEFAULT_BODY,
        default_text_body=DEFAULT_TEXT_BODY,
    )


@app.post("/preview")
@login_required
def preview():
    upload = request.files.get("spreadsheet")
    if not upload or not upload.filename:
        return render_template("index.html", error="Selecciona un archivo CSV o Excel.", default_body=DEFAULT_BODY), 400
    try:
        filename = secure_filename(upload.filename)
        headers, rows = read_spreadsheet(filename, upload.read())
        attach_enabled = request.form.get("attach_enabled") == "yes"
        attachments = {}
        uploaded_attachments = (
            [item for item in request.files.getlist("attachments") if item.filename]
            if attach_enabled
            else []
        )
        if len(uploaded_attachments) > MAX_ATTACHMENT_COUNT:
            raise ValueError(f"Solo se permiten hasta {MAX_ATTACHMENT_COUNT} archivos adjuntos por lote.")
        for uploaded in uploaded_attachments:
            if not uploaded.filename:
                continue
            attachment = validate_attachment(uploaded.filename, uploaded.read())
            key = attachment["filename"].casefold()
            if key in attachments:
                raise ValueError(f"Hay más de un adjunto con el nombre '{attachment['filename']}'.")
            attachments[key] = attachment
        body_template = request.form.get("body_template", DEFAULT_TEXT_BODY)
        campaign = {
            "rows": rows,
            "headers": headers,
            "source_name": filename,
            "attach_enabled": attach_enabled,
            "attachments": attachments,
            "subject_prefix": request.form.get("subject_prefix", "Trabajos de emergencia"),
            "suffix_mode": request.form.get("suffix_mode", "fixed"),
            "suffix_text": request.form.get("suffix_text", "23 y 24 de Setiembre"),
            "body_template": body_template,
            "body_format": "text",
            "owner_id": current_user.get_id(),
            "problems": validate_rows(rows, attach_enabled, set(attachments), body_template),
        }
    except (ValueError, UnicodeDecodeError, csv.Error, OSError) as error:
        return render_template("index.html", error=str(error), default_body=DEFAULT_BODY), 400
    campaign_id = secrets.token_urlsafe(18)
    persist_campaign(campaign_id, campaign, current_user.get_id())
    return redirect(url_for("preview_saved", campaign_id=campaign_id))


@app.get("/preview/<campaign_id>")
@login_required
def preview_saved(campaign_id):
    campaign, job = load_campaign(campaign_id, current_user.get_id())
    if campaign is None:
        abort(404)
    if job["status"] not in {"ready", "invalid"}:
        return redirect(url_for("history_detail", campaign_id=campaign_id))
    campaign["can_execute"] = job["status"] == "ready"
    return render_template("preview.html", campaign_id=campaign_id, campaign=campaign)


@app.get("/history")
@login_required
def history():
    jobs = get_db().execute(
        """SELECT id, source_name, created_at, completed_at, status, total_rows,
                  sent_count, error_count, dry_run
           FROM jobs WHERE user_id = ? ORDER BY created_at DESC LIMIT 100""",
        (int(current_user.get_id()),),
    ).fetchall()
    return render_template("history.html", jobs=jobs)


@app.get("/history/<campaign_id>")
@login_required
def history_detail(campaign_id):
    job = get_db().execute(
        "SELECT * FROM jobs WHERE id = ? AND user_id = ?",
        (campaign_id, int(current_user.get_id())),
    ).fetchone()
    if job is None:
        abort(404)
    rows = get_db().execute(
        """SELECT row_number AS row, recipient, attachment_name AS attachment,
                  status, detail FROM job_rows WHERE job_id = ? ORDER BY row_number""",
        (campaign_id,),
    ).fetchall()
    results = [
        {
            "row": row["row"],
            "recipient": row["recipient"],
            "attachment": row["attachment"],
            "ok": row["status"] == "sent",
            "pending": row["status"] in {"ready", "queued", "processing"},
            "status_label": {
                "ready": "Listo",
                "queued": "En cola",
                "processing": "Procesando",
                "sent": "Enviado",
                "failed": "Error",
                "invalid": "Inválido",
                "needs_review": "Revisión requerida",
            }.get(row["status"], row["status"]),
            "detail": row["detail"] or row["status"],
        }
        for row in rows
    ]
    return render_template("results.html", results=results, dry_run=bool(job["dry_run"]), job=job, historical=True)


@app.post("/send/<campaign_id>")
@login_required
def send(campaign_id):
    campaign, job = load_campaign(campaign_id, current_user.get_id())
    if campaign is None:
        abort(404)
    if job["status"] != "ready":
        return redirect(url_for("history_detail", campaign_id=campaign_id))
    if campaign["problems"]:
        return render_template("preview.html", campaign_id=campaign_id, campaign=campaign)
    if request.form.get("confirm_send") != "yes":
        return render_template("preview.html", campaign_id=campaign_id, campaign=campaign, error="Confirma que revisaste la vista previa.")
    password = request.form.get("smtp_password", "")
    if not password:
        return render_template("preview.html", campaign_id=campaign_id, campaign=campaign, error="Ingresa la clave SMTP para enviar.")
    database = get_db()
    database.execute("BEGIN IMMEDIATE")
    claimed = database.execute(
        "UPDATE jobs SET status = 'processing', dry_run = 0 WHERE id = ? AND user_id = ? AND status = 'ready'",
        (campaign_id, int(current_user.get_id())),
    ).rowcount
    if claimed:
        database.execute(
            "UPDATE job_rows SET status = 'processing' WHERE job_id = ? AND status = 'ready'",
            (campaign_id,),
        )
    database.commit()
    if not claimed:
        return redirect(url_for("history_detail", campaign_id=campaign_id))
    try:
        results = send_campaign(campaign, password)
        finish_campaign(campaign_id, results)
    except Exception as error:
        fail_campaign(campaign_id, f"No se pudo enviar el lote ({type(error).__name__}); revisa la clave SMTP y la conexión.")
    finally:
        password = ""
    return redirect(url_for("history_detail", campaign_id=campaign_id))


@app.errorhandler(413)
def too_large(_error):
    return render_template("index.html", error="El archivo supera el límite total de 25 MB.", default_body=DEFAULT_BODY), 413


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=False)
